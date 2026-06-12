"""XLSX export endpoints — downloadable spreadsheets with Destination Host column."""

import io

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook

from app.db.collections import col, RULE_REQUESTS, PHYSICAL_RULES, GROUP_CHANGE_REQUESTS, INGRESS_GROUPS
from app.services.export import build_xlsx_rows, resolve_destination_host

router = APIRouter(prefix="/api/export", tags=["Export"])


@router.get("/rules/{request_id}/manifest.xlsx")
async def export_rule_request_xlsx(
    request_id: str,
    dc_id: str | None = Query(None),
):
    """Download a rule request manifest as XLSX.

    The 'Destination Host' column shows VIP/endpoint references
    from ingress groups for documentation purposes.
    """
    req = await col(RULE_REQUESTS).find_one({"request_id": request_id})
    if not req:
        raise HTTPException(404, "Rule request not found")

    # Gather physical rules
    rule_query: dict = {"request_id": request_id}
    if dc_id:
        rule_query["$or"] = [{"src_dc": dc_id}, {"dst_dc": dc_id}]

    cursor = col(PHYSICAL_RULES).find(rule_query)
    physical_rules = await cursor.to_list(length=500)

    # Build manifest structure
    manifest = {
        "request_id": request_id,
        "requester": req.get("owner", ""),
        "owner_team": req.get("owner_team", ""),
        "environment": req.get("environment", ""),
        "status": req.get("status", ""),
        "rules": [
            {
                "rule_id": r.get("rule_id", ""),
                "vrf": "",
                "src_dc": r.get("src_dc", ""),
                "dst_dc": r.get("dst_dc", ""),
                "src_group": r.get("src_group_ref", ""),
                "dst_group": r.get("dst_group_ref", ""),
                "src_nh": r.get("src_nh", ""),
                "src_sz": r.get("src_sz", ""),
                "dst_nh": r.get("dst_nh", ""),
                "dst_sz": r.get("dst_sz", ""),
                "protocol": r.get("ports", "").split(" ")[0] if r.get("ports") else "",
                "ports": r.get("ports", ""),
                "action": r.get("action", ""),
                "lifecycle_status": r.get("lifecycle_status", ""),
            }
            for r in physical_rules
        ],
        "groups": [],
    }

    sheets = await build_xlsx_rows(manifest)

    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=(name or "Sheet")[:31])
        for row in rows:
            ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    suffix = f"-{dc_id}" if dc_id else ""
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{request_id}{suffix}-manifest.xlsx"'
        },
    )


@router.get("/groups/{group_name}/details.xlsx")
async def export_group_details_xlsx(group_name: str):
    """Download an ingress group's full details as XLSX.

    Includes members, VIP entries, and endpoint entries so the
    Destination Host information is available in spreadsheet form.
    """
    doc = await col(INGRESS_GROUPS).find_one({"name": group_name})
    if not doc:
        raise HTTPException(404, "Ingress group not found")

    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)

    # Members sheet
    ws_members = wb.create_sheet(title="Members")
    ws_members.append(["Type", "Value", "Description", "DC"])
    for m in doc.get("members", []):
        if isinstance(m, dict):
            ws_members.append([
                m.get("type", ""),
                m.get("value", ""),
                m.get("description", ""),
                m.get("dc_id", ""),
            ])

    # VIP Entries sheet (Destination Host references)
    ws_vips = wb.create_sheet(title="VIP Entries (Destination Host)")
    ws_vips.append(["VIP Address", "VIP Name", "Load Balancer", "Pool Members", "Description"])
    for vip in doc.get("vip_entries", []):
        ws_vips.append([
            vip.get("vip_address", ""),
            vip.get("vip_name", ""),
            vip.get("load_balancer", ""),
            ", ".join(vip.get("pool_members", [])),
            vip.get("description", ""),
        ])

    # Endpoint Entries sheet (Destination Host references)
    ws_endpoints = wb.create_sheet(title="Endpoints (Destination Host)")
    ws_endpoints.append(["Endpoint Name", "Type", "Resolved IPs", "Description"])
    for ep in doc.get("endpoint_entries", []):
        ws_endpoints.append([
            ep.get("endpoint_name", ""),
            ep.get("endpoint_type", "DNS"),
            ", ".join(ep.get("resolved_ips", [])),
            ep.get("description", ""),
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{group_name}-details.xlsx"'
        },
    )
